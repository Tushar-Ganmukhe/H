"""
Verify that stock was updated in Excel file
"""
import openpyxl
import pandas as pd

PRODUCT_FILE = r"C:\Users\hp\Downloads\products-export.xlsx"

print("Verifying Excel file...")
print("=" * 60)

# Check with openpyxl
print("\nUsing openpyxl:")
wb = openpyxl.load_workbook(PRODUCT_FILE)
ws = wb.active

# Find column indices
product_col = None
stock_col = None

for col_idx in range(1, ws.max_column + 1):
    header = ws.cell(row=1, column=col_idx).value
    if header and 'product name' in str(header).lower():
        product_col = col_idx
    if header and 'stock' in str(header).lower():
        stock_col = col_idx

# Find NORSAN
for row_idx in range(2, ws.max_row + 1):
    product_cell = ws.cell(row=row_idx, column=product_col)
    stock_cell = ws.cell(row=row_idx, column=stock_col)
    if product_cell.value and "NORSAN Omega-3 Total" in str(product_cell.value):
        print(f"Row {row_idx}: {product_cell.value}")
        print(f"  Stock: {stock_cell.value}")

wb.close()

# Check with pandas
print("\nUsing pandas:")
df = pd.read_excel(PRODUCT_FILE)
norsan_row = df[df['product name'].str.contains('NORSAN Omega-3 Total', na=False)]
if not norsan_row.empty:
    print(f"Found: {norsan_row['product name'].values[0]}")
    print(f"  Stock: {norsan_row['stock'].values[0]}")
else:
    print("Not found")

print("\n" + "=" * 60)
print("✓ Verification complete!")
