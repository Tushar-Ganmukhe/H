import openpyxl
import pandas as pd

PRODUCT_FILE = r"C:\Users\hp\Downloads\products-export.xlsx"

print("=" * 60)
print("INSPECT COLUMNS:")
print("=" * 60)

df = pd.read_excel(PRODUCT_FILE)
print(f"\nDataFrame columns: {df.columns.tolist()}")
print(f"\nFirst row:")
print(df.iloc[0])

print("\n" + "=" * 60)
print("OPENPYXL VIEW:")
print("=" * 60)
wb = openpyxl.load_workbook(PRODUCT_FILE)
ws = wb.active

print(f"\nHeaders from Excel (row 1):")
for col_idx in range(1, ws.max_column + 1):
    header = ws.cell(row=1, column=col_idx).value
    print(f"  Column {col_idx}: '{header}'")

# Find column indices for product and stock
product_col = None
stock_col = None

for col_idx in range(1, ws.max_column + 1):
    header = ws.cell(row=1, column=col_idx).value
    if header:
        header_lower = str(header).lower().strip()
        if 'product' in header_lower and 'name' in header_lower:
            product_col = col_idx
        if 'stock' in header_lower:
            stock_col = col_idx

print(f"\nIdentified columns - Product: {product_col}, Stock: {stock_col}")

if product_col and stock_col:
    print(f"\nFirst 15 products:")
    for row_idx in range(2, min(ws.max_row + 1, 17)):
        product_cell = ws.cell(row=row_idx, column=product_col)
        stock_cell = ws.cell(row=row_idx, column=stock_col)
        if product_cell.value:
            print(f"  Row {row_idx}: '{product_cell.value}' - Stock: {stock_cell.value}")
